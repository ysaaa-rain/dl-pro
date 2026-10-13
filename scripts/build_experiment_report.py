# -*- coding: utf-8 -*-
"""Build an editable Chinese report from verified experiment artifacts."""
import json, csv, statistics
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
S = json.loads((ROOT / "results/experiment/summary.json").read_text())
doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(2.2)
section.bottom_margin = Cm(2.2)
section.left_margin = Cm(2.3)
section.right_margin = Cm(2.3)


def font(run, size=12, bold=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    rf = run._element.get_or_add_rPr().get_or_add_rFonts()
    rf.set(qn("w:eastAsia"), "宋体")
    rf.set(qn("w:ascii"), "Times New Roman")
    rf.set(qn("w:hAnsi"), "Times New Roman")


for name, size in [("Normal", 12), ("Title", 18), ("Heading 1", 14), ("Heading 2", 12), ("Caption", 12)]:
    style = doc.styles[name]
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = name in ("Title", "Heading 1", "Heading 2")
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    rfonts = style._element.get_or_add_rPr().get_or_add_rFonts()
    for attribute in ("asciiTheme", "eastAsiaTheme", "hAnsiTheme", "cstheme"):
        rfonts.attrib.pop(qn("w:" + attribute), None)
    ppr = style._element.get_or_add_pPr()
    for border in list(ppr.findall(qn("w:pBdr"))):
        ppr.remove(border)
    style.paragraph_format.line_spacing = 1.25
    style.paragraph_format.space_before = Pt(3)
    style.paragraph_format.space_after = Pt(3)
    style.paragraph_format.first_line_indent = Pt(24 if name == "Normal" else 0)
    if name.startswith("Heading"):
        style.paragraph_format.keep_with_next = True


def para(text, style=None, indent=True):
    p = doc.add_paragraph(style=style)
    if not indent:
        p.paragraph_format.first_line_indent = Pt(0)
    font(p.add_run(text), 12)
    return p


def heading(text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    font(p.add_run(text), 14 if level == 1 else 12, True)
    return p


def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, width in enumerate(widths):
        t.columns[i].width = Cm(width)
    for c, text in zip(t.rows[0].cells, headers):
        c.text = str(text)
    for row in rows:
        cells = t.add_row().cells
        for c, text in zip(cells, row):
            c.text = str(text)
    for r, row in enumerate(t.rows):
        pr = row._tr.get_or_add_trPr()
        pr.append(OxmlElement("w:cantSplit"))
        if r == 0:
            pr.append(OxmlElement("w:tblHeader"))
        for c, width in zip(row.cells, widths):
            c.width = Cm(width)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcp = c._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                e = OxmlElement("w:" + edge)
                for key, value in [("val", "single"), ("sz", "4"), ("color", "D9D9D9")]:
                    e.set(qn("w:" + key), value)
                borders.append(e)
            tcp.append(borders)
            margin = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                e = OxmlElement("w:" + edge)
                e.set(qn("w:w"), "90")
                e.set(qn("w:type"), "dxa")
                margin.append(e)
            tcp.append(margin)
            if r == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "EEF0F2")
                tcp.append(shading)
            for p in c.paragraphs:
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    font(run, 12, r == 0)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def figure(path, caption, width=6.1):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(ROOT / path), width=Inches(width))
    shape._inline.docPr.set("descr", caption)
    p = para(caption, "Caption", False)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(title.add_run("MicroC 接触矩阵深度学习实验报告"), 18, True)
p = para("组号 ______    姓名与学号 ____________________", indent=False)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = para("结果日期 2026年10月13日", indent=False)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
para("本实验围绕已知结构识别、未注释结构候选发现、全基因组多条件可视化和接触矩阵超分辨重建展开。首轮计算已完成，生成了模型权重、逐位置预测、候选清单和图表。分类总体准确率为73.8%，但CHID召回率仍偏低；候选筛选的选择性尚未优于匹配预算的随机对照；超分辨改善了原生图像指标，却未全面改善O/E结构对比度。因此，本轮结论主要支持流程可运行和对照设计有效，不能宣称发现了新的生物学结构类型。")

heading("1 实验问题与数据")
para("输入为大肠杆菌MG1655的MicroC接触矩阵，参考序列为NC_000913.3，长度4,641,652 bp，原生分箱10 bp。完整集合含29个矩阵和19个样本条件标签。原有37℃野生型两个重复与NCBI GEO补充下载的27个矩阵均已核验可读；压缩原件和解压矩阵合计13.79 GB。")
para("已知结构共344条。课程表格的取整坐标与论文作者发布的BED/BEDPE文件逐条匹配，边界差异合计最大45 bp，保留原表并采用作者文件的零起点半开区间。基因轨道使用同一参考版本的4,516条基因记录，GFF的一起点闭区间已转换为零起点半开区间。作者代码仓库固定提交、文件来源和SHA256均已记录。")
with (ROOT / "data/annotations/structures.csv").open() as stream:
    annotations = list(csv.DictReader(stream))
length_rows = []
for kind in ("CHIN", "OPCID", "CHID"):
    lengths = [int(row["length_bp"]) for row in annotations if row["type"] == kind]
    length_rows.append([kind, len(lengths), f"{statistics.mean(lengths):.1f}", max(lengths)])
table(["类型", "结构数", "平均长度 bp", "最大长度 bp"], length_rows, [3.0, 2.5, 4.5, 4.5])
para("长度统计用于描述样本，不作为CNN输入。结构类别并非互斥的空间区域，例如CHID中可能包含CHIN；相邻窗口和两个重复不能当作独立样本随机拆分。")

heading("2 设计思路")
para("计算采用局部稀疏读取。全基因组float64稠密矩阵约需1.7 TB，不能直接在单机展开。读取局部接触像素后按目标分箱聚合，并按真实bin面积计算接触密度；环状起点窗口显式拼接，末端52 bp的100 bp分箱按实际宽度归一化。该聚合是对称接触图像的块平均，不等同于对唯一测序连接对重新计数。")
para("分类输入采用12.8 kb和25.6 kb两个尺度，均转换为128×128矩阵，分别对应100 bp和200 bp分箱。最大窗口覆盖所有已知结构。根据实际25.6 kb输入窗口的重叠连通关系形成56个位置组，先划分组，再扩展为两个生物学重复。距离期望和裁剪尺度仅在训练集合估计，主对角线及相邻一条对角线置零，以减轻线性距离背景。")
para("候选发现采用卷积自编码器学习未注释背景。整体强度保留、逐窗口强度归一化和计数噪声校正三种打分构成对照。最后一项在训练背景中拟合随平均接触计数变化的重建误差基线，再使用五个计数区间的稳健残差尺度进行校准。这样能够检验整体强度和稀疏计数噪声对异常分数的影响，而不是将高接触量直接认定为结构异常。")
para("超分辨以图像块均值构造40 bp输入，重建10 bp目标。模型在低分辨率空间学习卷积残差，经PixelShuffle放大4倍，控制原生大窗口的计算量。损失同时约束原生MSE和160 bp池化后的MSE。野生型rep1用于拟合，rep2仅用于评价；空间区块也独立划分。")
table(["模块", "主要参数", "选择与评价规则"], [
    ["分类CNN", "2通道 64维表征 35轮", "AdamW学习率0.001；验证macro-F1及NLL选检查点"],
    ["背景自编码器", "32维表征 30轮", "157个训练背景、51个验证背景；验证MSE选检查点"],
    ["超分辨CNN", "4倍放大 15轮", "64训练、12验证、32测试；验证MSE选检查点"],
], [3.0, 4.3, 7.2])

heading("3 已知结构识别与解释")
para("分类使用类别加权交叉熵，保留多数类、接触密度、注释长度和PCA线性分类器作为对照。注释长度基线是诊断性捷径：它依赖候选阶段未必可获得的标注信息，不能当作可部署的图像识别器。")
table(["首个空间留出对照", "准确率", "macro-F1"], [
    [r["method"], f'{r["accuracy"]:.3f}', f'{r["macro_f1"]:.3f}']
    for r in S["classification_baselines"]
], [7.0, 3.7, 3.8])
para("首个留出集合有65个位置，其中CHID仅2条，不能单独证明少数类识别稳定。三个训练种子的准确率范围为70.8%至76.9%，macro-F1范围为0.576至0.660。进一步保持同一训练配方，进行五折位置分组交叉验证，每个位置仅汇总两个重复的预测概率后评价一次。")
table(["类别", "样本数", "召回率", "F1"], [
    [c, int(S["per_class"][c]["support"]), f'{S["per_class"][c]["recall"]:.3f}', f'{S["per_class"][c]["f1-score"]:.3f}']
    for c in ["CHID", "CHIN", "OPCID"]
], [3.0, 3.0, 4.5, 4.0])
para(f'五折总体准确率为{S["classification_cv"]["accuracy"]:.3f}，macro-F1为{S["classification_cv"]["macro_f1"]:.3f}。按整个重叠组重采样所得95%区间分别为0.682至0.798和0.537至0.666。这些区间描述同一研究内部的位置组变动，不能替代外部数据验证。')
figure("results/classification_cv/confusion.png", "图1 五折折外预测混淆矩阵 每个结构位置只计一次", 4.2)
para("显著性图和等面积遮挡用于诊断模型关注区域。OPCID的中心遮挡平均降低真实类别概率0.290，角落遮挡反而增加概率；CHID的中心遮挡平均增加概率0.150，而角落遮挡降低概率0.244。该结果不支持模型对所有类别都主要关注中心结构。全量像素覆盖度审计发现两个重复各有8,428个零接触bin，占1.816%；将对应图像区域替换为中性背景后，65个留出位置中6个改变预测。零接触不能直接解释为映射失败或生物学缺失，其来源仍需核实。")

heading("4 候选发现与跨重复验证")
para("全基因组按6.4 kb步长扫描726个环状窗口。训练背景和验证背景按100 kb区块分离，并避开已知结构和区块边界。异常阈值仅用分类训练位置上的已知结构校准，目标为中心召回达到90%；候选去重要求中心间距至少12.8 kb。")
table(["打分对照", "与密度的相关性", "筛选窗口", "去重候选"], [
    [r["method"], f'{r["density_correlation"]:.3f}', r["selected"], r["retained"]]
    for r in S["discovery_controls"]
], [4.0, 4.5, 3.0, 3.0])
para("去除整体强度后，低计数窗口的噪声被放大，相关性从正值转为强负值。计数条件校准将相关性降到0.027，但保留窗口仍达645个，占扫描总数88.8%。当前方法的高召回依赖宽松候选预算，选择性仍然不足。")
table(["65个留出已知结构", "去重前中心召回", "去重后中心召回"], [
    ["当前检测器", "0.954", "0.738"],
    ["同预算随机对照 200次均值", "0.989", "0.814"],
], [6.5, 4.0, 4.0])
para("表中评价的是已知结构中心是否落入候选窗口，不是精确边界定位。随机对照优于当前打分，因此不能仅凭召回率高就认定检测器有效。表征的独立背景对照在清除窗口重叠后取得balanced accuracy 0.723和ROC-AUC 0.789，说明表征有一定区分信息，但异常打分仍需改进。")
para("对294个去重候选进行PCA降维和六簇K-means聚类。复现性使用去除主对角线和相邻对角线的O/E上三角相关；阈值至少为0.5，并与错配背景窗口的95%分位数比较。已知结构排除缓冲为200 bp。新簇要求不含已知区间，且至少有5个非重叠成员通过复现性门槛。")
para("得到37个未注释且跨重复可复现的窗口，但六个簇均含已知结构，符合预设新簇条件的数量为0。未注释仅说明当前344条注释中没有覆盖该窗口，不能据此认定它属于新的生物学类型。")
para("下面按跨重复相关性最高选择一个未注释窗口展示，选择规则不用于证明生物学新颖性。")
example = S["example_candidate"]
figure(f'results/discovery/heatmaps/{example["id"]}.png',
       f'图2 未注释窗口示例 {example["id"]} 中心{example["center"]} bp 相关性{example["repeat_correlation"]:.3f}', 6.1)

heading("5 多条件基因组轨道")
para("采用37℃野生型、44℃和rif25三个条件，每个条件两个重复。先按样本总有效接触对数归一化，再平均同条件重复。每个位点的局部强度定义为其与正负10 kb范围内其他位点的接触密度之和，排除自对角线；它不是整条染色体的总接触覆盖度。")
para("每张图包含条件曲线及灰色均值背景、平均信号、基因、已知结构四个面板，共享以bp为单位的横轴。固定10,000 bp区间生成465张图，最后一张截断到染色体末端，区间连续覆盖4,641,652 bp。基因使用蓝色方块，CHIN、OPCID和CHID使用橙色层级区间，重叠结构分层放置。")
figure("results/multitrack/tracks/0050000_0060000.png", "图3 50,000至60,000 bp区间的四轨展示")
figure("results/multitrack/tracks/0150000_0160000.png", "图4 150,000至160,000 bp区间的重叠结构与基因轨道")
para("这些曲线用于观察同一位置的条件差异。不同条件的测序深度已进行总量校正，但未执行ICE平衡，当前结果也没有独立机制实验支持，不能直接从曲线差异推断某蛋白的因果作用。")

heading("6 超分辨扩展实验")
para("训练窗口从先前的1.28 kb扩展到10.24 kb，40 bp输入恢复10 bp目标。训练只使用rep1的64个窗口，验证使用rep1的12个窗口；测试为独立区块中的16个位置、两个重复共32个窗口。rep2完全不参与模型拟合和检查点选择。强度尺度为训练计数精确99.5%分位数的log1p值，两个重复按总接触对数统一到rep1深度。")
table(["重复", "方法", "PSNR dB", "SSIM"], [
    [r["replicate"], {"bicubic": "双三次", "bilinear": "双线性", "subpixel_cnn": "CNN三种子均值"}[r["method"]],
     f'{r["psnr"]:.3f}', f'{r["ssim"]:.3f}'] for r in S["sr_native"]
], [2.0, 6.0, 3.4, 3.1])
para("所有三个种子的CNN均改善原生PSNR和SSIM。图像指标在统一裁剪的log强度空间计算，不能直接解释为真实10 bp接触被准确恢复。MPS加速设备及随机种子已记录，不保证不同机器间逐位一致。")
para("下游只选择完整25.6 kb上下文处于超分辨测试区块中的49个结构位置，含37条CHIN、8条OPCID、4条CHID。每个位置使用自身未参与拟合的折外分类器评价原始和重建图，并平均两个重复的概率。")
table(["下游方法", "分类准确率", "macro-F1", "与原始预测一致"], [
    [{"original": "原始", "clipped_target": "裁剪目标", "bilinear": "双线性", "bicubic": "双三次", "subpixel_cnn": "CNN"}[r["method"]],
     f'{r["accuracy"]:.3f}', f'{r["macro_f1"]:.3f}', f'{r["agreement_with_original_oof"]:.3f}']
    for r in S["sr_downstream"]
], [4.0, 3.0, 3.0, 4.5])
figure("results/report_figures/CHIN_restoration.png", "图5 完整CHIN上下文的O/E形状比较 青色框为源标注区间")
para("CNN的下游分类准确率从原始图的65.3%升至69.4%，但这只是49个位置和一个折外分类器的代理指标。与原始O/E图比较，双三次插值的平均SSIM为0.991，CNN为0.962；相关系数则分别约0.995和0.996。原生图像误差、区域对比度和分类保留出现不同结论，因此不能宣称结构保真度已全面优于插值。下一轮需加强对O/E结构对比度的损失约束，并检查空间误差的具体位置。")

heading("7 核验与局限")
para("八项测试通过，包含环状起点稀疏读取、末端短bin面积校正、强度倍增不改变形状输入、形状变化仍能被保留以及原有数据读取与泄漏检查。另核验了每个结构恰有一条折外预测、465张轨道的连续覆盖、rep2不参与超分辨拟合、重建数组有限且在规定范围内。")
para("当前macOS数值库在部分矩阵乘法中发出溢出等警告。独立非BLAS算术复核得到基线预测概率差小于8×10⁻¹⁵，目标梯度最大绝对值约7.2×10⁻⁵；PCA特征方程残差约10⁻¹⁵，K-means标签与独立距离计算一致。输出均有限，警告和复核记录保留，跨平台复跑仍需再次核验。")
para("本轮的主要局限是少数类CHID识别不稳定、候选异常打分选择性不足、未知背景可能包含未标注结构，以及超分辨的O/E对比度指标未获一致改善。全部评价来自同一研究，不代表外部数据泛化；几个随机种子也不是独立的生物学实验。没有得到满足本实验定义的新簇，因而没有新的生物学结构结论。")

heading("8 可复现记录与下一步")
para("输入文件来源与哈希位于样本清单，原始表格、版本化BED和GFF均保留。处理窗口分片、模型权重、逐位置指标、候选热图和全基因组轨道已保存。大体积原始矩阵和内存映射重建数组保留本地，可通过官方来源和脚本恢复。")
for command in [
    "python3 scripts/download_geo.py --include-wt",
    ".venv/bin/python scripts/prepare_geo.py",
    ".venv/bin/python scripts/resolve_reference.py",
    ".venv/bin/python scripts/prepare_experiments.py --reuse-cache",
    ".venv/bin/python scripts/train_classifier.py",
    ".venv/bin/python scripts/crossvalidate_classifier.py",
    ".venv/bin/python scripts/discover_structures.py",
    ".venv/bin/python scripts/multitrack.py",
    ".venv/bin/python scripts/sr_expanded.py",
]:
    p = para(command, indent=False)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
para("下一轮优先区分零接触区域与真实结构，检验有效像素掩码和更聚焦的多尺度表征；候选发现应在固定预算下比较形态方法、密度方法和随机方法；超分辨应明确权衡原生图像质量与O/E结构保真度。最终提交还需填写成员分工和贡献比例，并完成PPT讲解视频。")

heading("9 参考资料")
para("Gavrilov AA, Shamovsky I, et al. Elementary 3D organization of active and silenced E. coli genome. Nature, 2025, 645:1060–1070. https://www.nature.com/articles/s41586-025-09396-y", indent=False)
para("NCBI GEO GSE272159 MicroC contact matrices. https://www.ncbi.nlm.nih.gov/geo/query/acc.cgi?acc=GSE272159", indent=False)
para("论文作者的坐标与基因注释代码仓库。https://github.com/irzhegalova/ecoli_microc；固定提交与文件哈希见reference_resolution.json。", indent=False)
para("课程资料《基于深度学习的接触矩阵处理实践方案》，任务一至三以及选做任务五。", indent=False)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.paragraph_format.first_line_indent = Pt(0)
font(footer.add_run("第 "), 10)
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)
font(footer.add_run(" 页"), 10)
doc.core_properties.title = "MicroC 接触矩阵深度学习实验报告"
doc.core_properties.subject = "首轮真实数据实验与对照评价"
doc.core_properties.author = ""
dest = ROOT / "output/实验报告_第一轮计算结果.docx"
dest.parent.mkdir(exist_ok=True)
doc.save(dest)
print(dest)
