# -*- coding: utf-8 -*-
"""Build an editable Chinese report from verified experiment artifacts."""
import json, csv, statistics
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
S = json.loads((ROOT / "results/refinement/summary.json").read_text())
doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(2.2)
section.bottom_margin = Cm(2.2)
section.left_margin = Cm(2.3)
section.right_margin = Cm(2.3)


def font(run, size=12, bold=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    rf = run._element.get_or_add_rPr().get_or_add_rFonts()
    rf.set(qn("w:eastAsia"), "宋体")
    rf.set(qn("w:ascii"), "Times New Roman")
    rf.set(qn("w:hAnsi"), "Times New Roman")


for name, size in [("Normal", 12), ("Title", 18), ("Heading 1", 14), ("Heading 2", 12), ("Caption", 12)]:
    style = doc.styles[name]
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = name in ("Title", "Heading 1", "Heading 2")
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    rfonts = style._element.get_or_add_rPr().get_or_add_rFonts()
    for attribute in ("asciiTheme", "eastAsiaTheme", "hAnsiTheme", "cstheme"):
        rfonts.attrib.pop(qn("w:" + attribute), None)
    ppr = style._element.get_or_add_pPr()
    for border in list(ppr.findall(qn("w:pBdr"))):
        ppr.remove(border)
    style.paragraph_format.line_spacing = 1.25
    style.paragraph_format.space_before = Pt(3)
    style.paragraph_format.space_after = Pt(3)
    style.paragraph_format.first_line_indent = Pt(24 if name == "Normal" else 0)
    if name.startswith("Heading"):
        style.paragraph_format.keep_with_next = True


def para(text, style=None, indent=True):
    p = doc.add_paragraph(style=style)
    if not indent:
        p.paragraph_format.first_line_indent = Pt(0)
    font(p.add_run(text), 12)
    return p


def heading(text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    font(p.add_run(text), 14 if level == 1 else 12, True)
    return p


def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, width in enumerate(widths):
        t.columns[i].width = Cm(width)
    for c, text in zip(t.rows[0].cells, headers):
        c.text = str(text)
    for row in rows:
        cells = t.add_row().cells
        for c, text in zip(cells, row):
            c.text = str(text)
    for r, row in enumerate(t.rows):
        pr = row._tr.get_or_add_trPr()
        pr.append(OxmlElement("w:cantSplit"))
        if r == 0:
            pr.append(OxmlElement("w:tblHeader"))
        for c, width in zip(row.cells, widths):
            c.width = Cm(width)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcp = c._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                e = OxmlElement("w:" + edge)
                for key, value in [("val", "single"), ("sz", "4"), ("color", "D9D9D9")]:
                    e.set(qn("w:" + key), value)
                borders.append(e)
            tcp.append(borders)
            margin = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                e = OxmlElement("w:" + edge)
                e.set(qn("w:w"), "90")
                e.set(qn("w:type"), "dxa")
                margin.append(e)
            tcp.append(margin)
            if r == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "EEF0F2")
                tcp.append(shading)
            for p in c.paragraphs:
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    font(run, 12, r == 0)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def figure(path, caption, width=6.1):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(ROOT / path), width=Inches(width))
    shape._inline.docPr.set("descr", caption)
    p = para(caption, "Caption", False)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


title=doc.add_paragraph(style="Title");title.alignment=WD_ALIGN_PARAGRAPH.CENTER
font(title.add_run("MicroC 第二轮改进与对照报告"),18,True)
para("结果日期 2026年10月14日",indent=False)
para("本报告记录第一轮之后的候选选择性与超分辨结构保真改进。第一轮原件和全部失败对照保留。本轮固定预算下的候选召回点估计改善，计数一致性校正也缓解了超分辨结构失真；但新结构簇仍为零，新增超分辨训练没有优于对已有模型直接做计数校正。四项有结果仍不能等同于达到优秀水平。")
heading("1 数据与评价边界")
para("任务二与任务五继续使用WT 37℃两个重复。候选评分仅在rep1背景训练区拟合，验证标签选择三个形态汇总方式之一；rep2用于复现核验。超分辨沿用64训练、12验证、32测试窗口，训练和验证只含rep1，rep2只评价。所有第一轮测试位置已被观察，因此本轮是回顾性改进比较，不是新增的独立测试。")
para("本轮另新增WT、HNS_mut、FIS_mut三条件轨道，各两个重复。原来的WT、44℃、rif25图组保留。这次真正使用了老师补发清单中的HNS_mut和FIS_mut四个矩阵，正式实验输入去重共10个，剩余19个矩阵仍只做了数据审计。")
heading("2 固定预算的候选选择")
para("第一轮按训练已知结构90%召回选择阈值，候选预算过宽。本轮预先固定64个候选为主要预算，128个为辅助预算，两者均采用12.8 kb中心最小间距。新形态特征取两个窗口尺度内的O/E正向局部对比，以不同高斯平滑尺度提取连贯富集；近对角线和零接触行列不参与评分。背景计数条件回归仅在157个rep1训练背景窗口拟合。")
para("三个形态汇总方式只按验证集64候选的中心召回选择。选择结果为最大标准化局部对比，验证召回0.417。它是带验证标签选择的形态检测方法，不能描述为完全无监督新结构发现。密度、第一轮AE和500次随机选择都采用相同候选数及去重间距。")
rows=[]
for budget in [64,128]:
    for method,label in [(S['discovery']['score_choice'],'形态评分'),('first_round_AE','第一轮AE'),('density','密度评分')]:
        value=next(r['center_recall'] for r in S['discovery_test'] if r['budget']==budget and r['method']==method);rows.append([budget,label,f'{value:.3f}'])
    value=next(r['mean'] for r in S['random_test'] if r['budget']==budget);rows.append([budget,'随机500次均值',f'{value:.3f}'])
table(['候选预算','方法','原测试65个结构中心召回'],rows,[3,5,6.5])
figure('results/refinement/discovery_budget.png','图1 相同预算与去重间距下的召回对照')
para("64候选的形态召回为0.354，随机均值0.178，第一轮AE为0.185；128候选的形态召回为0.569，随机均值0.352。点估计改善，但原测试集只有12个重叠组。按组重采样得到64候选相对随机均值的召回差95%区间约为−0.054至0.319，包含零，不能断言已经稳健优于随机或跨数据泛化。形态评分与密度的Spearman相关仍为−0.267，密度混杂并未完全消失。")
para("64个主候选中有10个未注释且跨重复相关达标，满足至少5个复现成员并完全排除已知注释的新簇数量仍为0。背景训练PCA、六簇K-means、逐候选评分、重复相关和64张热图已保留。候选窗口不等于精确边界定位，未注释窗口也不等于新生物学类型。")
heading("3 超分辨的计数一致性与消融")
para("第一轮CNN改善裁剪log图像指标，却改变了接触块总量。本轮将高分辨率预测计数逐4×4块缩放，使均值等于实测40 bp输入；保留对称性并允许零计数块。原始40 bp计数在此约束中不作高强度裁剪，原生PSNR和SSIM仍在第一轮训练99.5%尺度裁剪空间评价。第一轮CNN仍接收原始裁剪输入，后处理使用实测未裁剪块均值。输入尾部信息差异已明确记录。")
para("新模型在相同划分上用三个种子训练15轮，增加100 bp计数域O/E损失，并按验证组合损失选权重。100 bp不是40 bp的整数倍，能检查跨块边缘的结构误差；200 bp为整数倍，其保留主要来自约束本身，不是学习成功的独立证据。另设直接校正双三次和直接校正第一轮CNN两种消融。")
rows=[]
for method,label in [('bicubic','双三次'),('projected_bicubic','计数校正双三次')]:
    v=next(r for r in S['sr_native'] if r['replicate']==1 and r['method']==method);rows.append([label,f"{v['psnr']:.3f}",f"{v['ssim']:.3f}"])
for method,label in [('first_round_cnn','第一轮CNN 种子42'),('projected_first_round_cnn','计数校正第一轮CNN')]:
    v=next(r for r in S['sr_projection_native'] if r['replicate']==1 and r['method']==method);rows.append([label,f"{v['psnr']:.3f}",f"{v['ssim']:.3f}"])
v=next(r for r in S['sr_native'] if r['replicate']==1 and r['method']=='subpixel_cnn');rows.append(['重新训练 三种子均值',f"{v['psnr']:.3f}",f"{v['ssim']:.3f}"])
table(['WT rep1方法','PSNR dB','SSIM'],rows,[8,3.25,3.25])
para("直接校正第一轮CNN的rep2 PSNR为21.640、SSIM为0.435，仍高于双三次21.531和0.410。rep1 PSNR由原CNN的22.205降到22.032，说明计数一致性也有图像误差代价；不能只展示改善的指标。新增训练的三种子均值为22.028和0.448，没有优于直接校正已有CNN。")
figure('results/refinement/sr_shape.png','图2 49个完整上下文的100 bp O/E评价')
para("同一49个位置、两个重复的100 bp O/E SSIM，双三次约0.903、第一轮CNN约0.870、计数校正第一轮CNN约0.915、重新训练模型约0.914。校正相对第一轮CNN的逐位置平均增量约0.045，位置重采样95%区间为0.0437至0.0464；该区间未按更大空间连通组或独立生物学样本重采样，不代表外部泛化。200 bp SSIM接近1主要由计数约束保证。")
para("重新训练及计数校正双三次的下游折外分类准确率为65.3%、macro-F1为0.543，与原始矩阵相同。第一轮CNN的69.4%下游准确率没有保留。因此本轮建议保留计数校正第一轮CNN作为结构一致性方案，同时保留原CNN作为图像误差方案，不宣称各项指标同时最优。真正10 bp块内细节仍缺乏独立证据。")
heading("4 老师补发条件与下一步")
para("多轨道脚本支持条件组合和独立输出目录，并按实际样本ID与来源哈希检查缓存归属，防止复用错误条件。新增465张WT、HNS_mut、FIS_mut轨道图连续覆盖4,641,652 bp，与第一轮图组分开保存。局部信号为深度归一化后的正负10 kb接触密度，重复先归一化再平均；无ICE平衡，不能据此推断蛋白直接形成结构。")
figure('results/multitrack_mutants/tracks/0150000_0160000.png','图3 WT HNS_mut FIS_mut的四轨展示')
para("新增计数守恒、梯度有限、对称性及形态评分强度不变测试，本轮共10项通过。本轮PCA仍有Accelerate警告，独立非BLAS特征方程及聚类距离复核记录已保存。第一轮任务一未重新训练，CHID召回34.6%仍是短板。尚未开发辅助标注工具或LLM辅助模块，也未完成PPT视频及成员贡献填写，不能把它们计作完成的加分项。")
heading("5 复跑与证据")
for command in ['.venv/bin/python scripts/discovery_refinement.py','.venv/bin/python scripts/sr_refinement.py','.venv/bin/python scripts/sr_projection_ablation.py','.venv/bin/python scripts/multitrack.py --conditions 37C HNS_mut FIS_mut --output results/multitrack_mutants']:
    p=para(command,indent=False);p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
para("结果分别位于results/discovery_refined、results/sr_refined、results/multitrack_mutants与results/refinement。原始数据、第一轮结果、负面消融和运行日志保留。模型权重、处理数据和图表可随仓库恢复，大重建数组留本地由脚本生成。本报告是本轮对照补充，不替代第一轮报告。")
footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.title='MicroC 第二轮改进与对照报告';doc.core_properties.author=''
dest=ROOT/'output/实验报告_第二轮改进对照.docx';doc.save(dest);print(dest)
